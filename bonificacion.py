"""
FindOut - Módulo Bonificación.

Cruza los "Puntos PM Conciliados" del DWH (op.FactNovedadesOperador, sumados
por operador) contra el archivo plano de bonificación cargado por el
usuario, y genera un resumen de coincidencias y alertas.
"""

import io
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill, Font

import db

COLUMNAS_ESPERADAS = [
    "Identificación", "Código", "Empleado", "Empresa", "Cargo", "Grupo",
    "Días Novedad", "Puntos PM Conciliados", "Valor Bono", "Vr Día",
    "Vr Punto", "Dcto Puntos", "Dcto Días", "A pagar",
]

COLOR_HEADER = "145A4F"
COLOR_OK = "C6EFCE"
COLOR_ALERTA = "F8CBAD"
COLOR_FALTANTE = "FFEB9C"


def _normalizar(texto) -> str:
    if pd.isna(texto):
        return ""
    return " ".join(str(texto).strip().upper().split())


def _conexion_manual():
    """
    Muestra en la barra lateral los parámetros de conexión al DWH, precargados
    desde .streamlit/secrets.toml cuando existen, con un botón para probar la
    conexión y ver el resultado al instante.
    """
    cfg = st.secrets.get("dwh", {}) if hasattr(st, "secrets") else {}

    with st.sidebar.expander("🔌 Conexión al DWH", expanded=True):
        st.caption("Precargado desde secrets.toml. Puedes editarlo aquí solo para esta sesión.")
        host = st.text_input("Host", value=str(cfg.get("host", "10.0.22.78")), key="man_host")
        port = st.text_input("Puerto", value=str(cfg.get("port", "5432")), key="man_port")
        database = st.text_input("Base de datos", value=str(cfg.get("database", "GRMDW")), key="man_db")
        user = st.text_input("Usuario", value=str(cfg.get("user", "richard.guevara")), key="man_user")
        password = st.text_input(
            "Contraseña", type="password", value=str(cfg.get("password", "")), key="man_pass"
        )

        col_probar, col_usar = st.columns(2)
        with col_probar:
            probar = st.button("🔎 Probar conexión", use_container_width=True)
        with col_usar:
            usar_manual = st.checkbox("Usar estos datos", key="man_usar")

        if probar:
            if not user or not password:
                st.error("Completa usuario y contraseña para probar la conexión.")
            else:
                with st.spinner("Conectando..."):
                    try:
                        engine_prueba = db.construir_engine(
                            host=host, port=port, database=database, user=user, password=password
                        )
                        db.probar_conexion(engine_prueba)
                        st.success(f"✅ Conexión exitosa a {database}@{host}:{port}")
                    except Exception as e:
                        st.error(f"❌ No se pudo conectar:\n\n{e}")

    if usar_manual and user and password:
        return db.construir_engine(host=host, port=port, database=database, user=user, password=password)
    return None


def render():
    st.markdown(
        '<div class="findout-header"><h1>Bonificación</h1>'
        '<p>Cruce de Puntos PM Conciliados — DWH (miproceso) vs. archivo plano</p></div>',
        unsafe_allow_html=True,
    )

    engine_manual = _conexion_manual()

    col1, col2 = st.columns(2)
    with col1:
        hoy = date.today()
        primer_dia_mes = hoy.replace(day=1)
        fecha_inicio = st.date_input("Fecha inicio novedades", value=primer_dia_mes)
    with col2:
        fecha_fin = st.date_input("Fecha fin novedades", value=hoy)

    solo_procede = st.checkbox("Considerar solo novedades con Procede = Sí", value=True)

    archivo = st.file_uploader("Cargar archivo plano de bonificación (Excel)", type=["xlsx", "xls"])

    df_archivo = None
    columna_llave = "Empleado"
    normalizar_texto = True

    if archivo is not None:
        try:
            df_archivo = pd.read_excel(archivo)
        except Exception as e:
            st.error(f"No se pudo leer el archivo: {e}")
            return

        faltantes = [c for c in COLUMNAS_ESPERADAS if c not in df_archivo.columns]
        if faltantes:
            st.warning(
                "El archivo no tiene todas las columnas esperadas. "
                f"Faltan: {', '.join(faltantes)}. Se continuará con las columnas disponibles."
            )

        st.success(f"Archivo cargado: {len(df_archivo)} registros.")
        with st.expander("Ver archivo cargado"):
            st.dataframe(df_archivo, use_container_width=True)

        colA, colB = st.columns(2)
        with colA:
            opciones_llave = [c for c in ["Empleado", "Identificación", "Código"] if c in df_archivo.columns]
            columna_llave = st.selectbox(
                "Columna del archivo para relacionar con 'Operador' del DWH",
                opciones_llave or df_archivo.columns.tolist(),
            )
        with colB:
            normalizar_texto = st.checkbox(
                "Normalizar texto para comparar (mayúsculas, sin espacios extra)", value=True
            )

    ejecutar = st.button(
        "🔍 Ejecutar validación", type="primary", use_container_width=True, disabled=archivo is None
    )

    if not ejecutar:
        return

    if fecha_inicio > fecha_fin:
        st.error("La fecha de inicio no puede ser posterior a la fecha fin.")
        return

    if "Puntos PM Conciliados" not in df_archivo.columns:
        st.error("El archivo debe contener la columna 'Puntos PM Conciliados' para poder cruzar la información.")
        return

    with st.spinner("Consultando novedades en el DWH..."):
        try:
            df_dwh_completo = db.obtener_novedades(fecha_inicio, fecha_fin, engine=engine_manual)
        except Exception as e:
            st.error(f"No fue posible consultar el DWH: {e}")
            return

    if df_dwh_completo.empty:
        st.warning("El DWH no devolvió novedades para el período seleccionado.")
        return

    # Inconsistencia a validar en los datos del DWH: novedades marcadas como
    # "Procede = No" que, sin embargo, tienen puntos PM conciliados mayores
    # a cero. Esto se revisa sobre TODAS las novedades del período, sin
    # importar el filtro de Procede elegido arriba para el cruce principal.
    df_anomalias_procede = df_dwh_completo[
        (~df_dwh_completo["EsProcede"]) & (df_dwh_completo["PuntosPmConciliados"] > 0)
    ].copy()

    df_dwh = df_dwh_completo[df_dwh_completo["EsProcede"]] if solo_procede else df_dwh_completo

    if df_dwh.empty:
        st.warning("No quedaron novedades del DWH después de aplicar el filtro de Procede.")
        return

    st.caption(
        f"Novedades obtenidas del DWH: {len(df_dwh_completo)} registros en el período, "
        f"{len(df_dwh)} consideradas para el cruce ({df_dwh['Operador'].nunique()} operadores)."
    )

    df_dwh_agrupado = (
        df_dwh.groupby("Operador", as_index=False)["PuntosPmConciliados"]
        .sum()
        .rename(columns={"PuntosPmConciliados": "Puntos DWH"})
    )

    df_cruce = df_archivo.copy()
    if normalizar_texto:
        df_cruce["_llave"] = df_cruce[columna_llave].apply(_normalizar)
        df_dwh_agrupado["_llave"] = df_dwh_agrupado["Operador"].apply(_normalizar)
    else:
        df_cruce["_llave"] = df_cruce[columna_llave]
        df_dwh_agrupado["_llave"] = df_dwh_agrupado["Operador"]

    resultado = df_cruce.merge(df_dwh_agrupado, on="_llave", how="outer", indicator=True)

    resultado["Puntos PM Conciliados"] = pd.to_numeric(resultado["Puntos PM Conciliados"], errors="coerce")
    resultado["Puntos DWH"] = pd.to_numeric(resultado["Puntos DWH"], errors="coerce")
    resultado["Diferencia"] = resultado["Puntos PM Conciliados"] - resultado["Puntos DWH"]

    def _estado(row):
        if row["_merge"] == "left_only":
            return "NO ENCONTRADO EN DWH"
        if row["_merge"] == "right_only":
            return "NO ENCONTRADO EN ARCHIVO"
        if pd.isna(row["Diferencia"]) or row["Diferencia"] != 0:
            return "ALERTA"
        return "OK"

    resultado["Estado"] = resultado.apply(_estado, axis=1)

    columnas_finales = [columna_llave, "Operador", "Puntos PM Conciliados", "Puntos DWH", "Diferencia", "Estado"]
    columnas_finales = [c for c in dict.fromkeys(columnas_finales) if c in resultado.columns]
    resultado_final = resultado[columnas_finales].sort_values("Estado").reset_index(drop=True)

    total = len(resultado_final)
    ok_n = int((resultado_final["Estado"] == "OK").sum())
    alerta_n = int((resultado_final["Estado"] == "ALERTA").sum())
    no_dwh_n = int((resultado_final["Estado"] == "NO ENCONTRADO EN DWH").sum())
    no_archivo_n = int((resultado_final["Estado"] == "NO ENCONTRADO EN ARCHIVO").sum())

    st.markdown("### Consistencia de novedades en el DWH")
    if df_anomalias_procede.empty:
        st.success("No se encontraron novedades con Procede = No que tengan puntos PM conciliados mayores a cero.")
    else:
        st.error(
            f"⚠️ Se encontraron {len(df_anomalias_procede)} novedades marcadas como **Procede = No** "
            "que igualmente tienen puntos PM conciliados mayores a cero. Esto no debería ocurrir y debe revisarse."
        )
        with st.expander(f"Ver detalle de la inconsistencia ({len(df_anomalias_procede)} registros)", expanded=True):
            st.dataframe(
                df_anomalias_procede[
                    ["Fecha", "Operador", "TipoNovedad", "PmGrupo", "DetalleNovedad", "PuntosPmConciliados", "Procede"]
                ].sort_values(["Operador", "Fecha"]),
                use_container_width=True,
            )

    st.markdown("### Resumen del cruce")
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Total registros", total)
    m2.metric("Coinciden (OK)", ok_n)
    m3.metric("Con diferencias", alerta_n)
    m4.metric("No en DWH", no_dwh_n)
    m5.metric("No en archivo", no_archivo_n)

    if alerta_n or no_dwh_n or no_archivo_n:
        st.warning(
            f"Se identificaron {alerta_n} diferencias de puntos, {no_dwh_n} empleados del archivo "
            f"sin novedades en el DWH y {no_archivo_n} operadores del DWH que no están en el archivo."
        )
    else:
        st.success("Todos los registros coinciden exactamente entre el DWH y el archivo cargado.")

    def _resaltar(row):
        color = {
            "OK": "background-color: #E4F5EF",
            "ALERTA": "background-color: #FBE4E1",
            "NO ENCONTRADO EN DWH": "background-color: #FDF3D9",
            "NO ENCONTRADO EN ARCHIVO": "background-color: #FDF3D9",
        }.get(row["Estado"], "")
        return [color] * len(row)

    st.markdown("### Detalle del cruce")
    st.dataframe(resultado_final.style.apply(_resaltar, axis=1), use_container_width=True, height=420)

    with st.expander(f"🔴 Alertas y diferencias ({alerta_n + no_dwh_n + no_archivo_n})"):
        st.dataframe(resultado_final[resultado_final["Estado"] != "OK"], use_container_width=True)

    excel_bytes = _generar_reporte_excel(resultado_final, df_anomalias_procede)
    st.download_button(
        "⬇️ Descargar reporte de validación (Excel)",
        data=excel_bytes,
        file_name=f"FindOut_Bonificacion_{fecha_inicio}_{fecha_fin}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )


def _generar_reporte_excel(df: pd.DataFrame, df_anomalias_procede: pd.DataFrame = None) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cruce Bonificacion")
        ws = writer.sheets["Cruce Bonificacion"]

        fill_por_estado = {
            "OK": PatternFill(start_color=COLOR_OK, end_color=COLOR_OK, fill_type="solid"),
            "ALERTA": PatternFill(start_color=COLOR_ALERTA, end_color=COLOR_ALERTA, fill_type="solid"),
            "NO ENCONTRADO EN DWH": PatternFill(start_color=COLOR_FALTANTE, end_color=COLOR_FALTANTE, fill_type="solid"),
            "NO ENCONTRADO EN ARCHIVO": PatternFill(start_color=COLOR_FALTANTE, end_color=COLOR_FALTANTE, fill_type="solid"),
        }

        encabezados = [c.value for c in ws[1]]
        col_estado = encabezados.index("Estado") + 1 if "Estado" in encabezados else None

        if col_estado:
            for fila in range(2, ws.max_row + 1):
                estado = ws.cell(row=fila, column=col_estado).value
                fill = fill_por_estado.get(estado)
                if fill:
                    for celda in ws[fila]:
                        celda.fill = fill

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")

        for columna in ws.columns:
            longitud = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
            ws.column_dimensions[columna[0].column_letter].width = min(longitud + 2, 40)

        if df_anomalias_procede is not None and not df_anomalias_procede.empty:
            columnas_anom = [
                "Fecha", "Operador", "TipoNovedad", "PmGrupo", "DetalleNovedad", "PuntosPmConciliados", "Procede"
            ]
            columnas_anom = [c for c in columnas_anom if c in df_anomalias_procede.columns]
            df_anomalias_procede[columnas_anom].to_excel(
                writer, index=False, sheet_name="Procede No con Puntos"
            )
            ws2 = writer.sheets["Procede No con Puntos"]

            fill_alerta = PatternFill(start_color=COLOR_ALERTA, end_color=COLOR_ALERTA, fill_type="solid")
            for fila in range(2, ws2.max_row + 1):
                for celda in ws2[fila]:
                    celda.fill = fill_alerta

            for cell in ws2[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")

            for columna in ws2.columns:
                longitud = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
                ws2.column_dimensions[columna[0].column_letter].width = min(longitud + 2, 40)

    buffer.seek(0)
    return buffer.getvalue()
